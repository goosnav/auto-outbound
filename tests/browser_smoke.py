"""Exercise an isolated running desktop release, including actual downloads."""

import argparse
import csv
import io
import json
from pathlib import Path

from openpyxl import load_workbook
from playwright.sync_api import expect, sync_playwright


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--url', required=True)
    parser.add_argument('--output', type=Path, default=Path('build/browser-smoke'))
    args = parser.parse_args()
    args.output.mkdir(parents=True, exist_ok=True)
    with sync_playwright() as playwright:
        browser = playwright.chromium.launch(headless=True)
        page = browser.new_page(viewport={'width': 1440, 'height': 1000}, device_scale_factor=1)
        errors = []
        page.on('pageerror', lambda error: errors.append(str(error)))
        page.goto(args.url)
        page.wait_for_load_state('networkidle')
        page.get_by_role('button', name='New list', exact=True).click()
        page.get_by_role('button', name='Create list', exact=True).click()
        assert page.locator('#campaign-dialog').is_visible()
        page.get_by_label('Name', exact=True).fill('Precision shops')
        page.get_by_label('Business type', exact=True).fill('Independent precision machine shops')
        page.get_by_label('Location', exact=True).fill('Southern California')
        page.locator('#new-provider').select_option('demo')
        page.get_by_role('button', name='Create list', exact=True).click()
        expect(page.get_by_role('heading', name='Precision shops', exact=True)).to_be_visible()
        assert page.locator('#run-stage').input_value() == 'report'
        expect(page.locator('[data-action="export"][data-format="csv"]')).to_be_disabled()
        page.locator('.heading-actions [data-action="run"]').click()
        expect(page.locator('.metric-value').nth(1)).to_have_text('4', timeout=45000)
        expect(page.locator('.heading-actions [data-action="run"]')).to_be_enabled(timeout=45000)
        campaign_id = page.locator('.campaign-link.active').get_attribute('data-id')
        detail = page.request.get(args.url + '/api/campaigns/' + campaign_id).json()
        assert detail['counts']['drafted'] == 0
        expect(page.locator('#job-region')).to_be_empty()
        expect(page.locator('[data-action="export"][data-scope="all"][data-format="csv"]')).to_be_enabled()
        page.locator('#toast-region').evaluate('(element) => element.replaceChildren()')
        page.evaluate('window.scrollTo(0, 0)')
        page.screenshot(path=str(args.output / 'leads.png'), full_page=True)
        page.get_by_role('searchbox').fill('no such company')
        expect(page.get_by_role('heading', name='No matching leads.')).to_be_visible()
        page.get_by_role('button', name='Reset filters').click()
        page.locator('.company-name').first.click()
        expect(page.get_by_role('dialog', name=page.locator('#lead-dialog-title').inner_text())).to_be_visible()
        page.get_by_role('tab', name='Evidence', exact=False).click()
        assert page.locator('.evidence-card').count() > 0
        page.screenshot(path=str(args.output / 'evidence.png'), full_page=True)
        page.get_by_role('button', name='Close lead review').click()
        page.get_by_role('button', name='Report', exact=True).click()
        expect(page.get_by_role('heading', name='Precision shops', exact=True)).to_be_visible()
        assert page.get_by_role('radio', name='All companies', exact=False).is_checked()
        for name, filename in [('Download CSV', 'sample-report.csv'), ('Excel', 'sample-report.xlsx')]:
            with page.expect_download() as download:
                page.get_by_role('button', name=name, exact=True).click()
            download.value.save_as(args.output / filename)
        rows = list(csv.DictReader(io.StringIO((args.output / 'sample-report.csv').read_text('utf-8-sig'))))
        assert len(rows) == 4 and all(row['source_urls'] for row in rows)
        assert not any(row['email_1_body'] for row in rows)
        book = load_workbook(args.output / 'sample-report.xlsx')
        assert book.sheetnames == ['Summary', 'Leads', 'Read me']
        assert book['Summary']['B6'].value == 4
        assert book['Leads'].max_row == 5
        page.locator('#toast-region').evaluate('(element) => element.replaceChildren()')
        page.evaluate('window.scrollTo(0, 0)')
        page.screenshot(path=str(args.output / 'report.png'), full_page=True)
        for width in (390, 320, 768):
            page.set_viewport_size({'width': width, 'height': 900})
            assert page.evaluate('document.documentElement.scrollWidth <= innerWidth'), f'Report overflows at {width}px'
            page.get_by_role('button', name='Leads', exact=True).click()
            assert page.evaluate('document.documentElement.scrollWidth <= innerWidth'), f'List overflows at {width}px'
            page.screenshot(path=str(args.output / f'mobile-{width}.png'), full_page=True)
            page.get_by_role('button', name='Setup', exact=True).click()
            assert page.evaluate('document.documentElement.scrollWidth <= innerWidth'), f'Brief overflows at {width}px'
            page.get_by_role('button', name='Report', exact=True).click()
        page.set_viewport_size({'width': 720, 'height': 500})
        assert page.evaluate('document.documentElement.scrollWidth <= innerWidth')
        assert not errors, errors
        (args.output / 'result.json').write_text(json.dumps({'result': 'PASS', 'csv_rows': len(rows), 'sheets': book.sheetnames, 'widths': [1440, 768, 390, 320], 'page_errors': errors}, indent=2) + '\n')
        browser.close()
        print('PASS: no-offer report run, evidence, filters, real CSV/XLSX downloads, 320/390/768/1440px layouts; no browser errors.')


if __name__ == '__main__':
    main()
