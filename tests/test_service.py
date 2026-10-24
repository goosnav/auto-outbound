import csv
import io
from datetime import UTC, datetime, timedelta

import pytest
from auto_outbound import demo
from auto_outbound.config import AppError
from auto_outbound.models import LeadSeed, now, public_url
from auto_outbound.pipeline import Jobs
from auto_outbound.prompts import queries
from auto_outbound.service import Service
from openpyxl import load_workbook


@pytest.fixture
def service(tmp_path):
    return Service(tmp_path)


@pytest.fixture
def prepared(service):
    c = service.create_campaign({"name": "Fixture", "icp": {"business_type": "Machine shops"}, "offer": "RFQ review queue", "sender": {"name": "Test Sender", "company": "Fixture Studio", "email": "sender@example.com", "postal_address": "1 Example Street, Los Angeles, CA"}})
    seed = LeadSeed.model_validate(demo.discovery_result(1, [])["companies"][0])
    lead, _ = service.upsert_seed(c["id"], seed)
    lead.update(demo.research_result(lead))
    lead.update(research_status="complete", status="drafted", drafts=demo.sequence_result(lead)["drafts"])
    service.store.save("leads", lead)
    return c, lead


def make_ready(service, lead):
    return service.update_lead(lead["id"], {"email_status": "verified", "verification_source": "Fixture verifier", "verification_at": now(), "status": "approved"})


def test_demo_full_workflow_persists_and_is_never_sendable(service):
    campaign = service.create_demo()
    job = Jobs(service).create(campaign["id"], background=False)
    assert job["status"] == "completed"
    assert job["agent_calls"] == 9
    reopened = Service(service.store.root).campaign_detail(campaign["id"])
    assert reopened["counts"]["drafted"] == 4
    assert all(len(lead["drafts"]) == 3 and not lead["qa_issues"] for lead in reopened["leads"])
    assert all(not lead["ready_to_send"] for lead in reopened["leads"])
    with pytest.raises(AppError, match="No eligible"):
        service.export(campaign["id"], scope="approved")
    with pytest.raises(AppError, match="No eligible"):
        service.export(campaign["id"], scope="callable")


def test_source_and_www_domains_dedupe_without_overwriting_user_contact(service, prepared):
    campaign, lead = prepared
    seed = LeadSeed(company="Changed title", website=lead["website"].replace("https://", "https://www.") + "/", email="other@example.com")
    merged, created = service.upsert_seed(campaign["id"], seed)
    assert not created
    assert merged["email"] == lead["email"]
    assert len(service.store.list("leads", campaign["id"])) == 1


def test_mx_is_not_mailbox_verification(service, prepared):
    campaign, lead = prepared
    result = service.update_lead(lead["id"], {"email_status": "mx_found", "status": "approved"})
    assert not result["ready_to_send"]
    assert "Mailbox verification needed" in result["export_blockers"]


def test_ready_export_requires_review_verifier_sender_and_actual_sources(service, prepared):
    campaign, lead = prepared
    ready = make_ready(service, lead)
    assert ready["ready_to_send"]
    raw = service.export(campaign["id"], scope="approved")
    rows = list(csv.DictReader(io.StringIO(raw.decode("utf-8-sig"))))
    assert len(rows) == 1 and rows[0]["email_3_subject"]
    assert rows[0]["footer"].endswith('reply "no thanks".')
    service.update_campaign(campaign["id"], {"sender": {"postal_address": ""}})
    assert not service.campaign_detail(campaign["id"])["leads"][0]["ready_to_send"]


@pytest.mark.parametrize("field,value", [("drafts", None), ("contact_name", "Changed Person"), ("phone", "+1 202-555-0101")])
def test_content_edits_revoke_approval(service, prepared, field, value):
    campaign, lead = prepared
    make_ready(service, lead)
    if field == "drafts":
        value = lead["drafts"]
        value[0]["subject"] = "Quote intake"
    updated = service.update_lead(lead["id"], {field: value})
    assert updated["status"] == "drafted"
    assert not updated["ready_to_send"]


def test_changing_email_invalidates_mailbox_evidence(service, prepared):
    _, lead = prepared
    make_ready(service, lead)
    updated = service.update_lead(lead["id"], {"email": "new@example.com"})
    assert updated["email_status"] == "unknown"
    assert updated["verification_source"] == updated["verification_at"] == updated["email_source"] == ""


def test_suppression_wins_over_approval_reruns_and_imports(service, prepared):
    campaign, lead = prepared
    make_ready(service, lead)
    service.suppress({"kind": "domain", "value": lead["domain"], "reason": "Existing customer"})
    with pytest.raises(AppError, match="suppressed"):
        service.update_lead(lead["id"], {"status": "approved"})
    service.upsert_seed(campaign["id"], LeadSeed.model_validate({key: lead[key] for key in LeadSeed.model_fields}))
    assert service.store.get("leads", lead["id"])["status"] == "suppressed"
    with pytest.raises(AppError, match="No eligible"):
        service.export(campaign["id"], scope="approved")


def test_verifier_import_marks_invalid_and_catch_all_without_fake_verified(service, prepared):
    campaign, lead = prepared
    header = "email,status,verification_source,verification_at\n"
    stamp = now()
    assert service.import_verification(campaign["id"], "verify.csv", (header + f"{lead['email']},catch-all,fixture,{stamp}\n").encode())["updated"] == 1
    assert service.store.get("leads", lead["id"])["email_status"] == "catch_all"
    service.import_verification(campaign["id"], "verify.csv", (header + f"{lead['email']},invalid,fixture,{stamp}\n").encode())
    assert service.store.get("leads", lead["id"])["status"] == "suppressed"


def test_future_or_missing_verification_dates_are_not_accepted(service, prepared):
    campaign, lead = prepared
    future = (datetime.now(UTC) + timedelta(days=1)).isoformat()
    with pytest.raises(AppError, match="actual verification date"):
        service.update_lead(lead["id"], {"email_status": "verified", "verification_source": "test", "verification_at": future})
    raw = f"email,status,verification_source\n{lead['email']},valid,test\n".encode()
    assert service.import_verification(campaign["id"], "result.csv", raw)["skipped"] == 1


def test_old_verification_expires_for_sending(service, prepared):
    _, lead = prepared
    old = (datetime.now(UTC) - timedelta(days=31)).isoformat()
    result = service.update_lead(lead["id"], {"email_status": "verified", "verification_source": "test", "verification_at": old, "status": "approved"})
    assert not result["ready_to_send"]


@pytest.mark.parametrize("outcome", ["bounced", "unsubscribed", "not_interested"])
def test_outcomes_create_global_suppression(service, prepared, outcome):
    campaign, lead = prepared
    service.import_outcomes(campaign["id"], "outcomes.csv", f"email,outcome\n{lead['email']},{outcome}\n".encode())
    assert service.is_suppressed(lead)


def test_prior_conversation_in_another_campaign_blocks_repeat_contact(service, prepared):
    campaign, lead = prepared
    service.import_outcomes(campaign["id"], "outcomes.csv", f"email,outcome\n{lead['email']},replied\n".encode())
    other = service.create_campaign({"name": "Another campaign"})
    duplicate, _ = service.upsert_seed(other["id"], LeadSeed.model_validate({key: lead[key] for key in LeadSeed.model_fields}))
    view = service.decorate(other, duplicate)
    assert any("another campaign" in reason for reason in view["export_blockers"])
    with pytest.raises(AppError, match="No eligible"):
        service.export(other["id"], scope="callable")


def test_shared_mailbox_updates_all_accounts_but_exports_one_message(service, prepared):
    campaign, first = prepared
    seed = LeadSeed.model_validate({key: first[key] for key in LeadSeed.model_fields} | {"company": "Second business", "website": "https://second.example"})
    second, _ = service.upsert_seed(campaign["id"], seed)
    second.update(drafts=first["drafts"], research_status="complete", fit_score=first["fit_score"] - 1, status="drafted", evidence=first["evidence"])
    service.store.save("leads", second)
    result = service.import_verification(campaign["id"], "checks.csv", f"email,status,provider,checked_at\n{first['email']},valid,Fixture verifier,{now()}\n".encode())
    assert result["updated"] == 2
    for lead in (first, second):
        service.update_lead(lead["id"], {"status": "approved"})
    exported = list(csv.DictReader(io.StringIO(service.export(campaign["id"], scope="approved").decode("utf-8-sig"))))
    assert len(exported) == 1
    assert exported[0]["company"] == first["company"]
    result = service.import_outcomes(campaign["id"], "results.csv", f"email,outcome\n{first['email']},replied\n".encode())
    assert result == {"updated": 2, "skipped": 0}
    assert all(lead["outcome"] == "replied" for lead in service.store.list("leads", campaign["id"]))
    with pytest.raises(AppError, match="No eligible"):
        service.export(campaign["id"], scope="approved")


def test_upload_dedupes_and_invalidates_drafts(service, prepared):
    campaign, lead = prepared
    make_ready(service, lead)
    one = service.upload_document(campaign["id"], "../../offer.md", b"New supporting offer details")
    two = service.upload_document(campaign["id"], "offer.md", b"New supporting offer details")
    assert one["id"] == two["id"]
    assert one["name"] == "offer.md"
    assert service.store.get("leads", lead["id"])["drafts"] == []
    service.delete_document(campaign["id"], one["id"])
    assert not service.store.documents(campaign["id"])


def test_csv_import_maps_standard_headers_and_reports_bad_rows(service):
    campaign = service.create_campaign({"name": "CSV"})
    data = b'Company Name,Company Website,First Name,Last Name,Work Email\nFixture,https://fixture.example,Jo,Example,jo@fixture.example\nMissing site,,,,\n'
    result = service.import_leads(campaign["id"], "contacts.csv", data)
    assert result["imported"] == 1 and result["rejected"] == 1
    lead = service.store.list("leads", campaign["id"])[0]
    assert lead["contact_name"] == "Jo Example"
    assert lead["email_status"] == "unknown"


def test_spreadsheet_preserves_multiline_drafts_and_sources(service, prepared):
    campaign, lead = prepared
    book = load_workbook(io.BytesIO(service.export(campaign["id"], "xlsx")))
    sheet = book["Leads"]
    headers = [cell.value for cell in sheet[1]]
    assert "\n" in sheet.cell(2, headers.index("email_1_body") + 1).value
    assert sheet.cell(2, headers.index("email_source") + 1).value == lead["email_source"]


def test_discovery_plan_has_unique_queries_above_eight(service):
    campaign = service.create_campaign({"name": "Large search", "icp": {"business_type": "machine shops"}, "settings": {"discovery_queries": 40}})
    plan = queries(campaign)
    assert len(plan) == len(set(plan)) == 40


@pytest.mark.parametrize("url", ["javascript:alert(1)", "file:///etc/passwd", "https://user@example.com", "https://:password@example.com"])
def test_nonweb_or_credentialed_urls_rejected(url):
    with pytest.raises(ValueError):
        public_url(url)


def test_report_workflow_needs_no_offer_or_sender_and_exports_research(service):
    campaign = service.create_demo()
    campaign = service.update_campaign(campaign["id"], {"offer": "", "sender": {"name": "", "company": "", "email": "", "postal_address": ""}})
    job = Jobs(service).create(campaign["id"], "report", background=False)
    assert job["status"] == "completed"
    assert job["stage_index"] == job["stage_count"] == 3
    detail = service.campaign_detail(campaign["id"])
    assert detail["counts"]["researched"] == 4
    assert detail["counts"]["drafted"] == 0
    assert all(not lead["drafts"] for lead in detail["leads"])
    rows = list(csv.DictReader(io.StringIO(service.export(campaign["id"]).decode("utf-8-sig"))))
    assert len(rows) == 4 and all(row["source_urls"] for row in rows)
    workbook = load_workbook(io.BytesIO(service.export(campaign["id"], "xlsx")))
    assert workbook.sheetnames == ["Summary", "Leads", "Read me"]
    assert workbook["Summary"]["B6"].value == 4
    assert workbook["Leads"].max_row == 5
    rerun = Jobs(service).create(campaign["id"], "report", background=False)
    assert rerun["status"] == "completed"
    assert service.campaign_detail(campaign["id"])["counts"]["total"] == 4


def test_changed_brief_is_marked_in_report_without_losing_saved_evidence(service):
    campaign = service.create_demo()
    Jobs(service).create(campaign["id"], "report", background=False)
    before = service.campaign_detail(campaign["id"])
    assert before["counts"]["researched"] == 4
    service.update_campaign(campaign["id"], {"icp": {"location": "A different region"}})
    after = service.campaign_detail(campaign["id"])
    assert after["counts"]["researched"] == 0
    assert all(not lead["research_current"] and lead["evidence"] for lead in after["leads"])
    rows = list(csv.DictReader(io.StringIO(service.export(campaign["id"]).decode("utf-8-sig"))))
    assert all(row["research_current"] == "no" for row in rows)
    Jobs(service).create(campaign["id"], "research", background=False)
    assert service.campaign_detail(campaign["id"])["counts"]["researched"] == 4
