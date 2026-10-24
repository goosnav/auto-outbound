"""Bounded document imports and spreadsheet-safe exports."""

from __future__ import annotations

import csv
import io
import json
import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from pypdf import PdfReader

from .config import AppError
from .quality import sender_footer

MAX_FILE = 8 * 1024 * 1024
MAX_ROWS = 10000


def check_file(raw: bytes):
    if not raw:
        raise AppError("FILE-EMPTY", "Choose a file with content.")
    if len(raw) > MAX_FILE:
        raise AppError("FILE-SIZE", "Files must be 8 MB or smaller.", 413)


def check_zip(raw: bytes):
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            items = archive.infolist()
            if len(items) > 2000 or sum(i.file_size for i in items) > 40 * 1024 * 1024:
                raise AppError("FILE-SIZE", "The expanded document is too large.")
            if any(i.flag_bits & 1 for i in items):
                raise AppError("FILE-ENCRYPTED", "Use an unencrypted document.")
    except zipfile.BadZipFile as exc:
        raise AppError("FILE-FORMAT", "This document is not a valid Office file.") from exc


def document_text(name: str, raw: bytes) -> str:
    check_file(raw)
    suffix = Path(name).suffix.lower()
    try:
        if suffix in {".txt", ".md", ".csv"}:
            text = raw.decode("utf-8-sig")
        elif suffix == ".pdf":
            reader = PdfReader(io.BytesIO(raw))
            if reader.is_encrypted:
                raise AppError("FILE-ENCRYPTED", "Use an unencrypted PDF.")
            if len(reader.pages) > 80:
                raise AppError("FILE-SIZE", "Use a PDF with 80 pages or fewer.")
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        elif suffix == ".docx":
            check_zip(raw)
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                xml = archive.read("word/document.xml")
            root = ElementTree.fromstring(xml)
            text = "\n".join(
                "".join(p.itertext())
                for p in root.iter(
                    "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p"
                )
            )
        else:
            raise AppError("FILE-TYPE", "Use TXT, Markdown, PDF, DOCX, or CSV for your offer.")
    except AppError:
        raise
    except Exception as exc:
        raise AppError(
            "FILE-READ", "The document could not be read. Try a plain-text export."
        ) from exc
    text = text.replace("\x00", "").strip()
    if not text:
        raise AppError("FILE-NO-TEXT", "No text found. Scanned PDFs need OCR before upload.")
    if len(text) > 60000:
        raise AppError("FILE-SIZE", "Use a shorter document (up to 60,000 text characters).")
    return text


def header(value) -> str:
    return re.sub(r"[^a-z0-9]+", "_", str(value or "").strip().lower()).strip("_")


def read_rows(name: str, raw: bytes) -> list[dict]:
    check_file(raw)
    suffix = Path(name).suffix.lower()
    try:
        if suffix == ".csv":
            text = raw.decode("utf-8-sig")
            try:
                dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t")
            except csv.Error:
                dialect = csv.excel
            reader = csv.reader(io.StringIO(text, newline=""), dialect)
            rows = list(_bounded(reader))
        elif suffix == ".xlsx":
            check_zip(raw)
            book = load_workbook(io.BytesIO(raw), read_only=True, data_only=False)
            try:
                rows = list(_bounded(book.active.iter_rows(values_only=True)))
            finally:
                book.close()
        else:
            raise AppError("FILE-TYPE", "Import a CSV or XLSX file.")
        if len(rows) < 2:
            raise AppError("FILE-ROWS", "Include a header and at least one data row.")
        keys = [header(k) for k in rows[0]]
        if any(not key for key in keys) or len(keys) != len(set(keys)):
            raise AppError("FILE-HEADERS", "Use unique, nonempty column names.")
        return [
            {
                key: str(value).strip() if value is not None else ""
                for key, value in zip(keys, row, strict=False)
                if key
            }
            for row in rows[1:]
            if any(value is not None and str(value).strip() for value in row)
        ]
    except AppError:
        raise
    except Exception as exc:
        raise AppError(
            "FILE-READ", "Could not read this spreadsheet. Try a UTF-8 CSV export."
        ) from exc


def _bounded(rows):
    for index, row in enumerate(rows):
        if index > MAX_ROWS or len(row) > 150:
            raise AppError("FILE-SIZE", "Imports support up to 10,000 rows and 150 columns.")
        if any(len(str(value or "")) > 60000 for value in row):
            raise AppError("FILE-SIZE", "A cell is too large.")
        yield row


def safe_cell(value):
    if value is None:
        return ""
    if isinstance(value, (dict, list)):
        value = json.dumps(value, ensure_ascii=False)
    if isinstance(value, bool):
        return "yes" if value else "no"
    if isinstance(value, str):
        value = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", value)
        if value.lstrip().startswith(("=", "+", "-", "@")) or value.startswith(("\t", "\r")):
            return "'" + value
    return value


BASE_COLUMNS = [
    "company",
    "website",
    "domain",
    "location",
    "industry",
    "size",
    "description",
    "contact_name",
    "contact_role",
    "email",
    "phone",
    "linkedin_url",
    "social_urls",
    "fit_score",
    "fit_reasons",
    "email_status",
    "email_source",
    "phone_source",
    "contact_source",
    "email_checked_at",
    "verification_source",
    "verification_at",
    "research_status",
    "research_current",
    "status",
    "ready_to_send",
    "export_blockers",
    "personalization",
    "pain_hypothesis",
    "research_summary",
    "risks",
    "review_notes",
    "outcome",
    "outcome_at",
]
EXPORT_COLUMNS = (
    BASE_COLUMNS
    + [
        f"email_{step}_{field}"
        for step in (1, 2, 3)
        for field in ("subject", "body", "delay_days", "source_urls")
    ]
    + [
        "sender_name",
        "sender_company",
        "sender_email",
        "sender_postal_address",
        "footer",
        "source_urls",
        "evidence",
        "created_at",
        "updated_at",
        "lead_id",
        "campaign_id",
    ]
)


def export_rows(campaign: dict, leads: list[dict]) -> list[dict]:
    result = []
    for lead in leads:
        row = {key: lead.get(key, "") for key in BASE_COLUMNS}
        row["export_blockers"] = "; ".join(lead.get("export_blockers", []))
        for draft in lead.get("drafts", []):
            for field in ("subject", "body", "delay_days", "source_urls"):
                row[f"email_{draft['step']}_{field}"] = draft[field]
        for key in ("name", "company", "email", "postal_address"):
            row[f"sender_{key}"] = campaign["sender"][key]
        row["footer"] = sender_footer(campaign["sender"])
        for key in ("source_urls", "evidence", "created_at", "updated_at", "campaign_id"):
            row[key] = lead.get(key, "")
        row["lead_id"] = lead["id"]
        result.append(row)
    return result


def spreadsheet(campaign: dict, leads: list[dict], format: str) -> bytes:
    rows = export_rows(campaign, leads)
    if format == "csv":
        buffer = io.StringIO(newline="")
        writer = csv.DictWriter(buffer, fieldnames=EXPORT_COLUMNS)
        writer.writeheader()
        writer.writerows(
            {key: safe_cell(row.get(key, "")) for key in EXPORT_COLUMNS} for row in rows
        )
        return buffer.getvalue().encode("utf-8-sig")
    if format != "xlsx":
        raise AppError("EXPORT-FORMAT", "Choose CSV or XLSX.")
    book = Workbook()
    summary = book.active
    summary.title = "Summary"
    for label, value in (
        ("AUTO OUTBOUND", "Research report"),
        ("List", campaign["name"]),
        ("Data", "FICTIONAL DEMO — DO NOT CONTACT" if campaign["demo"] else "Saved company research"),
        ("Business type", campaign.get("icp", {}).get("business_type", "")),
        ("Location", campaign.get("icp", {}).get("location", "")),
        ("Companies in this export", len(leads)),
        ("Current research", sum(lead.get("research_current", lead.get("research_status") == "complete") for lead in leads)),
        ("With email", sum(bool(lead.get("email")) for lead in leads)),
        ("Mailbox verified", sum(lead.get("email_status") == "verified" for lead in leads)),
        ("Report contents", "Open the Leads sheet for contacts, research, sources, and remaining gaps."),
        ("Contact status", "A public email or MX record does not prove mailbox delivery. Review exclusions before contact."),
    ):
        summary.append([label, safe_cell(value)])
    summary.column_dimensions["A"].width = 28
    summary.column_dimensions["B"].width = 80
    for row in summary:
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        row[0].font = Font(bold=True)
    for cell in summary[1]:
        cell.fill = PatternFill("solid", fgColor="1746C7")
        cell.font = Font(color="FFFFFF", bold=True)
    summary.freeze_panes = "B2"
    sheet = book.create_sheet("Leads")
    sheet.append(EXPORT_COLUMNS)
    for row in rows:
        sheet.append([safe_cell(row.get(key, "")) for key in EXPORT_COLUMNS])
    sheet.freeze_panes = "D2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="172320")
        cell.font = Font(color="FFFFFF", bold=True)
        sheet.column_dimensions[cell.column_letter].width = 25
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    notes = book.create_sheet("Read me")
    notes.append(["Campaign", campaign["name"]])
    notes.append(
        [
            "Data",
            "FICTIONAL DEMO — DO NOT CONTACT"
            if campaign["demo"]
            else "Research and drafts; no messages have been sent.",
        ]
    )
    notes.append(
        ["Verification", "MX-found means mail routing exists. It does not verify a mailbox."]
    )
    notes.append(
        [
            "Sending",
            "Use ready_to_send=yes only. Add footer to every body. Configure replies, bounces and unsubscribe handling in your sending tool.",
        ]
    )
    notes.append(
        [
            "Evidence",
            "Public sources are agent-reported unless a check is stated. Review personalization and claims.",
        ]
    )
    notes.append(
        ["Delay", "Suggested days from first send: 0, 4, 9. Stop on reply, bounce, or opt-out."]
    )
    notes.column_dimensions["A"].width = 20
    notes.column_dimensions["B"].width = 110
    for row in notes:
        for cell in row:
            cell.value = safe_cell(cell.value)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()
